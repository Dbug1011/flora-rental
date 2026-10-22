"""Build sheets-template/Flora_Booking_Sheet_Template.xlsx from the tabs Code.gs creates.
Usage: node scripts/build-template.js > /tmp/t.json && python3 scripts/build-template.py /tmp/t.json"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

data = json.load(open(sys.argv[1]))
wb = Workbook()
wb.remove(wb.active)
head_fill = PatternFill('solid', fgColor='F7D5E5')
head_font = Font(bold=True, color='4B203D')
widths = {'Catalog': 22, 'Bookings': 18, 'Booking Items': 18, 'Blocked Dates': 18, 'Settings': 30}
for name, rows in data.items():
    ws = wb.create_sheet(name)
    ncol = len(rows[0])
    for r in rows:
        ws.append([(v if v != '' else None) for v in (r + [''] * (ncol - len(r)))[:ncol]])
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
    ws.freeze_panes = 'A2'
    for i in range(1, ncol + 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths[name]
    # text format for codes, dates, times and mobile numbers
    text_cols = {'Booking code', 'Submission time', 'Mobile number', 'Pickup date', 'Pickup time', 'Return date', 'Return deadline',
                 'Hold expiry', 'Request key', 'Last updated', 'Item ID', 'Start date', 'End date', 'Key', 'Value'}
    for i, h in enumerate(rows[0], start=1):
        if h in text_cols:
            for r in range(2, 501):
                ws.cell(r, i).number_format = '@'
    if name == 'Settings':
        ws.column_dimensions['B'].width = 70
        ws.column_dimensions['C'].width = 70
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
    if name == 'Catalog':
        col = rows[0].index('Published') + 1
        dv = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{ws.cell(2, col).column_letter}2:{ws.cell(2, col).column_letter}500')
wb.save('sheets-template/Flora_Booking_Sheet_Template.xlsx')
print('saved')
